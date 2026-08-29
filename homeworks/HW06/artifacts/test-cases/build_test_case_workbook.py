#!/usr/bin/env python3
"""Build the HW06 Excel test-case workbook from the Markdown artifacts.

Sources (all read-only, nothing is invented here):

  artifacts/test-cases/<fr>/audit/audited-master-test-cases-v2.md  final 7-column suite
  artifacts/test-cases/<fr>/audit/extended-test-cases-v2.md        student additions (9 columns)
  artifacts/test-cases/<fr>/audit/audit-log-v2.md                  VALID/INVALID/INCOMPLETE labels
  artifacts/newman/<fr>/test-execution.md                          Newman per-test-case results
  BUG_REPORT.md                                                    bug summary table

Output: artifacts/test-cases/<StudentID>_HW06_TestCases.xlsx
  Sheet 1  Test Summary       counts per API, audit labels, bug list
  Sheet 2  FR-02 Login
  Sheet 3  FR-08 Checkout
  Sheet 4  FR-15 Product CRUD

Usage:
  python3 build_test_case_workbook.py [-o OUTPUT.xlsx]
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

STUDENT_ID = "23127216"
STUDENT_NAME = "Le Hoang Lam"

HW_ROOT = Path(__file__).resolve().parent.parent.parent  # homeworks/HW06
TEST_CASES_DIR = HW_ROOT / "artifacts" / "test-cases"
NEWMAN_DIR = HW_ROOT / "artifacts" / "newman"
BUG_REPORT = HW_ROOT / "BUG_REPORT.md"

FR_CONFIG = [
    {
        "key": "fr-02-login",
        "sheet": "FR-02 Login",
        "title": "FR-02 Login & Account Lockout",
        "pool": "A",
        "endpoint": "POST /api/login",
    },
    {
        "key": "fr-08-checkout",
        "sheet": "FR-08 Checkout",
        "title": "FR-08 Checkout",
        "pool": "B",
        "endpoint": "POST /api/checkout",
    },
    {
        "key": "fr-15-product-crud",
        "sheet": "FR-15 Product CRUD",
        "title": "FR-15 Product Management (Admin CRUD)",
        "pool": "C",
        "endpoint": "POST /api/products, PUT /api/products/:id, DELETE /api/products/:id",
    },
]

CASE_COLUMNS = [
    ("TC ID", 10),
    ("Origin", 15),
    ("Technique", 20),
    ("Precondition", 46),
    ("Input / action", 52),
    ("Expected result", 60),
    ("Oracle", 40),
    ("Trace", 34),
    ("Audit label", 13),
    ("Audit verdict", 14),
    ("Student fix", 46),
    ("Miss category", 16),
    ("Why the AI missed it", 52),
    ("Execution result", 15),
    ("Assertions", 16),
    ("Failure evidence", 46),
]

# ---------------------------------------------------------------- styling ---

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
SECTION_FONT = Font(bold=True, size=11, color="1F3864")
BODY_FONT = Font(size=10)
MONO_FONT = Font(size=10, name="Consolas")

FILL_PASS = PatternFill("solid", fgColor="C6EFCE")
FILL_FAIL = PatternFill("solid", fgColor="FFC7CE")
FILL_NEUTRAL = PatternFill("solid", fgColor="FFEB9C")
FILL_STUDENT = PatternFill("solid", fgColor="DDEBF7")
FILL_TOTAL = PatternFill("solid", fgColor="D9D9D9")

FONT_PASS = Font(size=10, bold=True, color="006100")
FONT_FAIL = Font(size=10, bold=True, color="9C0006")
FONT_NEUTRAL = Font(size=10, bold=True, color="9C6500")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TOP_WRAP = Alignment(vertical="top", wrap_text=True)
TOP_CENTER = Alignment(vertical="top", horizontal="center", wrap_text=True)


# ---------------------------------------------------------------- parsing ---


def clean(text: str) -> str:
    """Normalise a Markdown table cell into plain spreadsheet text."""
    text = text.replace("\\|", "|").replace("<br>", "\n").replace("<br/>", "\n")
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)  # keep link text only
    text = text.replace("`", "")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def split_row(line: str) -> list[str]:
    """Split one Markdown table row, honouring escaped pipes."""
    body = line.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|"):
        body = body[:-1]
    parts = re.split(r"(?<!\\)\|", body)
    return [clean(p) for p in parts]


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip()))


def read_tables(path: Path, first_header: str) -> list[dict[str, str]]:
    """Read every Markdown table in `path` whose first header cell matches.

    Rows with a column count different from the header are skipped and
    reported on stderr rather than silently mangled.
    """
    if not path.exists():
        raise FileNotFoundError(path)

    rows: list[dict[str, str]] = []
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            header = split_row(line)
            if header and header[0] == first_header and i + 1 < len(lines) and is_separator(lines[i + 1]):
                i += 2
                while i < len(lines) and lines[i].startswith("|"):
                    cells = split_row(lines[i])
                    if len(cells) == len(header):
                        rows.append(dict(zip(header, cells)))
                    elif any(cells):
                        print(
                            f"  ! {path.name}:{i + 1} column mismatch "
                            f"({len(cells)} vs {len(header)}) — row skipped",
                            file=sys.stderr,
                        )
                    i += 1
                continue
        i += 1
    return rows


TC_RE = re.compile(r"^TC-(\d+)([a-z]?)$", re.IGNORECASE)


def tc_sort_key(tc_id: str) -> tuple[int, str]:
    m = TC_RE.match(tc_id)
    if not m:
        return (10**6, tc_id)
    return (int(m.group(1)), m.group(2))


def tc_base(tc_id: str) -> str:
    """TC-38a -> TC-38 (sub-cases are one logical test case)."""
    m = TC_RE.match(tc_id)
    return f"TC-{int(m.group(1)):02d}" if m else tc_id


def normalise_tc(tc_id: str) -> str:
    m = TC_RE.match(tc_id)
    return f"TC-{int(m.group(1)):02d}{m.group(2).lower()}" if m else tc_id


def parse_execution(path: Path) -> tuple[dict[str, str], dict[str, dict]]:
    """Parse a Newman test-execution.md into (run metadata, per-TC results)."""
    meta: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("## "):  # metadata bullets only, before the first section
            break
        m = re.match(r"^- ([^:]+): (.+)$", line.strip())
        if m:
            meta.setdefault(m.group(1).strip(), clean(m.group(2)))

    results: dict[str, dict] = {}
    for row in read_tables(path, "Test case"):
        raw_id = row["Test case"]
        if not TC_RE.match(raw_id):  # FIXTURE / TEARDOWN helper rows
            meta.setdefault("_non_tc_rows", "")
            meta["_non_tc_rows"] += f"{raw_id}={row['Result']} "
            continue

        key = tc_base(raw_id)
        assertions = row.get("Assertions", "")
        n_pass = int(m.group(1)) if (m := re.search(r"(\d+) pass", assertions)) else 0
        n_fail = int(m.group(1)) if (m := re.search(r"(\d+) fail", assertions)) else 0

        entry = results.setdefault(
            key, {"result": "PASS", "pass": 0, "fail": 0, "evidence": [], "subcases": []}
        )
        entry["pass"] += n_pass
        entry["fail"] += n_fail
        entry["subcases"].append(normalise_tc(raw_id))
        if row.get("Result", "").upper() != "PASS":
            entry["result"] = row["Result"].upper()
        if row.get("Evidence"):
            entry["evidence"].append(f"{normalise_tc(raw_id)}: {row['Evidence']}")

    return meta, results


def parse_bugs(path: Path) -> list[dict[str, str]]:
    return read_tables(path, "Bug ID") or read_tables(path, "ID")


def collect_fr(cfg: dict) -> dict:
    """Assemble every row of one FR sheet from its four source documents."""
    audit_dir = TEST_CASES_DIR / cfg["key"] / "audit"

    suite = read_tables(audit_dir / "audited-master-test-cases-v2.md", "TC ID")
    extended = read_tables(audit_dir / "extended-test-cases-v2.md", "TC ID")
    audit_log = read_tables(audit_dir / "audit-log-v2.md", "TC ID")
    exec_meta, exec_results = parse_execution(NEWMAN_DIR / cfg["key"] / "test-execution.md")

    ext_by_id = {normalise_tc(r["TC ID"]): r for r in extended}
    log_by_id = {normalise_tc(r["TC ID"]): r for r in audit_log}

    rows = []
    for entry in suite:
        tc = normalise_tc(entry["TC ID"])
        ext = ext_by_id.get(tc, {})
        log = log_by_id.get(tc, {})
        res = exec_results.get(tc, {})

        assertions = ""
        if res:
            assertions = f"{res['pass']} pass / {res['fail']} fail"
            if len(res["subcases"]) > 1:
                assertions += f" ({len(res['subcases'])} sub-cases)"

        rows.append(
            {
                "TC ID": tc,
                "Origin": "Student-added" if ext else "AI-generated",
                "Technique": entry.get("Technique", ""),
                "Precondition": entry.get("Precondition", ""),
                "Input / action": entry.get("Input / action", ""),
                "Expected result": entry.get("Expected result", ""),
                "Oracle": entry.get("Oracle", ""),
                "Trace": entry.get("Trace", ""),
                "Audit label": log.get("New label", "N/A (student-added)" if ext else ""),
                "Audit verdict": log.get("Verdict", ""),
                "Student fix": log.get("Student fix", ""),
                "Miss category": ext.get("Miss category", ""),
                "Why the AI missed it": ext.get("Why the AI missed it", ""),
                "Execution result": res.get("result", "NOT RUN"),
                "Assertions": assertions,
                "Failure evidence": "; ".join(res.get("evidence", [])),
            }
        )

    rows.sort(key=lambda r: tc_sort_key(r["TC ID"]))

    unmatched = sorted(set(exec_results) - {r["TC ID"] for r in rows}, key=tc_sort_key)
    if unmatched:
        print(f"  ! {cfg['key']}: executed but not in the suite: {', '.join(unmatched)}", file=sys.stderr)

    # Collection Runner iterations: one CSV row per sub-case, fixtures excluded.
    runner_rows = sum(len(v["subcases"]) for v in exec_results.values())

    return {"cfg": cfg, "rows": rows, "exec_meta": exec_meta, "runner_rows": runner_rows}


# ---------------------------------------------------------------- writing ---


def style_header(ws: Worksheet, row_idx: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = TOP_CENTER
        cell.border = BORDER


def result_style(value: str) -> tuple[PatternFill | None, Font]:
    upper = value.upper()
    if upper == "PASS":
        return FILL_PASS, FONT_PASS
    if upper == "FAIL":
        return FILL_FAIL, FONT_FAIL
    if upper in {"NOT RUN", ""}:
        return None, BODY_FONT
    return FILL_NEUTRAL, FONT_NEUTRAL


def label_style(value: str) -> tuple[PatternFill | None, Font]:
    upper = value.upper()
    if upper == "VALID":
        return FILL_PASS, FONT_PASS
    if upper == "INVALID":
        return FILL_FAIL, FONT_FAIL
    if upper == "INCOMPLETE":
        return FILL_NEUTRAL, FONT_NEUTRAL
    return None, BODY_FONT


def write_case_sheet(wb: Workbook, data: dict) -> None:
    cfg = data["cfg"]
    ws = wb.create_sheet(cfg["sheet"])

    ws.cell(row=1, column=1, value=f"{cfg['title']} — {cfg['endpoint']}").font = TITLE_FONT
    meta = data["exec_meta"]
    ws.cell(
        row=2,
        column=1,
        value=(
            f"Suite: {len(data['rows'])} test cases "
            f"({sum(1 for r in data['rows'] if r['Origin'] == 'AI-generated')} AI-generated, "
            f"{sum(1 for r in data['rows'] if r['Origin'] == 'Student-added')} student-added)  |  "
            f"Newman assertions: {meta.get('Assertions', 'n/a')}  |  "
            f"Newman failures: {meta.get('Newman failures', 'n/a')}"
        ),
    ).font = BODY_FONT

    header_row = 4
    for idx, (name, width) in enumerate(CASE_COLUMNS, start=1):
        ws.cell(row=header_row, column=idx, value=name)
        ws.column_dimensions[get_column_letter(idx)].width = width
    style_header(ws, header_row, len(CASE_COLUMNS))

    for r_off, row in enumerate(data["rows"], start=header_row + 1):
        for c_idx, (name, _) in enumerate(CASE_COLUMNS, start=1):
            cell = ws.cell(row=r_off, column=c_idx, value=row[name])
            cell.alignment = TOP_WRAP
            cell.border = BORDER
            cell.font = BODY_FONT

            if name == "TC ID":
                cell.font = Font(size=10, bold=True)
                cell.alignment = TOP_CENTER
            elif name == "Origin":
                cell.alignment = TOP_CENTER
                if row[name] == "Student-added":
                    cell.fill = FILL_STUDENT
                    cell.font = Font(size=10, bold=True, color="1F3864")
            elif name in {"Execution result", "Audit label"}:
                fill, font = (
                    result_style(row[name]) if name == "Execution result" else label_style(row[name])
                )
                if fill:
                    cell.fill = fill
                cell.font = font
                cell.alignment = TOP_CENTER
            elif name == "Assertions":
                cell.alignment = TOP_CENTER

    last_row = header_row + len(data["rows"])
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(CASE_COLUMNS))}{last_row}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    ws.sheet_view.zoomScale = 90


def write_table(ws: Worksheet, start_row: int, headers: list[str], rows: list[list], widths=None) -> int:
    """Write one block table; returns the row index after the block."""
    for idx, name in enumerate(headers, start=1):
        ws.cell(row=start_row, column=idx, value=name)
    style_header(ws, start_row, len(headers))
    if widths:
        for idx, width in enumerate(widths, start=1):
            current = ws.column_dimensions[get_column_letter(idx)].width or 0
            ws.column_dimensions[get_column_letter(idx)].width = max(current, width)

    for r_off, row in enumerate(rows, start=start_row + 1):
        is_total = str(row[0]).upper().startswith("TOTAL")
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_off, column=c_idx, value=value)
            cell.border = BORDER
            cell.alignment = TOP_WRAP if c_idx <= 2 else TOP_CENTER
            cell.font = Font(size=10, bold=True) if is_total else BODY_FONT
            if is_total:
                cell.fill = FILL_TOTAL
    return start_row + len(rows) + 2


def write_summary_sheet(wb: Workbook, collected: list[dict], bugs: list[dict]) -> None:
    ws = wb.create_sheet("Test Summary", 0)

    ws.cell(row=1, column=1, value="HW06 — API Testing: Test Case Summary").font = TITLE_FONT
    info = [
        ("Student", f"{STUDENT_NAME} ({STUDENT_ID})"),
        ("System under test", "EShop (backend API, http://localhost:3000)"),
        ("Tooling", "Postman + Newman; header X-Student-Id: " + STUDENT_ID),
        ("Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("Source of truth", "audited-master-test-cases-v2.md, extended-test-cases-v2.md, audit-log-v2.md, newman/test-execution.md"),
    ]
    row = 3
    for key, value in info:
        ws.cell(row=row, column=1, value=key).font = Font(size=10, bold=True)
        ws.cell(row=row, column=2, value=value).font = BODY_FONT
        row += 1
    row += 1

    # --- per-API counts ---
    ws.cell(row=row, column=1, value="1. Test cases per API").font = SECTION_FONT
    row += 1
    headers = [
        "API", "Pool", "Endpoint(s)", "AI-generated", "Student-added", "Total",
        "Executed", "Passed", "Failed", "Pass rate", "Runner rows", "Newman assertions", "Bugs",
    ]
    widths = [30, 6, 46, 14, 14, 8, 10, 9, 9, 10, 12, 18, 8]
    body, totals = [], Counter()
    for data in collected:
        rows_ = data["rows"]
        ai = sum(1 for r in rows_ if r["Origin"] == "AI-generated")
        student = sum(1 for r in rows_ if r["Origin"] == "Student-added")
        executed = sum(1 for r in rows_ if r["Execution result"] != "NOT RUN")
        passed = sum(1 for r in rows_ if r["Execution result"] == "PASS")
        failed = sum(1 for r in rows_ if r["Execution result"] == "FAIL")
        fr_code = data["cfg"]["key"].split("-")[1]
        n_bugs = sum(1 for b in bugs if f"FR{fr_code}" in b.get("Bug ID", b.get("ID", "")))

        totals.update(
            {"ai": ai, "student": student, "total": len(rows_), "executed": executed,
             "passed": passed, "failed": failed, "bugs": n_bugs,
             "runner": data["runner_rows"]}
        )
        body.append([
            data["cfg"]["title"], data["cfg"]["pool"], data["cfg"]["endpoint"],
            ai, student, len(rows_), executed, passed, failed,
            f"{passed / executed:.0%}" if executed else "n/a",
            data["runner_rows"], data["exec_meta"].get("Assertions", "n/a"), n_bugs,
        ])

    body.append([
        "TOTAL (3 APIs)", "", "", totals["ai"], totals["student"], totals["total"],
        totals["executed"], totals["passed"], totals["failed"],
        f"{totals['passed'] / totals['executed']:.0%}" if totals["executed"] else "n/a",
        totals["runner"], "", totals["bugs"],
    ])
    row = write_table(ws, row, headers, body, widths)

    # --- audit labels ---
    ws.cell(row=row, column=1, value="2. Audit outcome of the AI-generated cases (Pass 2 labels)").font = SECTION_FONT
    row += 1
    label_rows = []
    label_totals = Counter()
    for data in collected:
        counts = Counter(
            r["Audit label"].upper() for r in data["rows"] if r["Origin"] == "AI-generated"
        )
        corrected = sum(
            1 for r in data["rows"]
            if r["Origin"] == "AI-generated" and r["Student fix"] and r["Student fix"].lower() != "none"
        )
        label_totals.update(counts)
        label_totals["corrected"] += corrected
        label_rows.append([
            data["cfg"]["title"], "", "",
            counts.get("VALID", 0), counts.get("INCOMPLETE", 0), counts.get("INVALID", 0), corrected,
        ])
    label_rows.append([
        "TOTAL", "", "",
        label_totals["VALID"], label_totals["INCOMPLETE"], label_totals["INVALID"],
        label_totals["corrected"],
    ])
    row = write_table(
        ws, row,
        ["API", "", "", "VALID", "INCOMPLETE", "INVALID", "Corrected by student"],
        label_rows, [30, 6, 46, 12, 14, 12, 20],
    )

    # --- why the AI missed the student additions ---
    ws.cell(row=row, column=1, value="3. Student-added cases by miss category").font = SECTION_FONT
    row += 1
    categories = ["prompt quality", "model limitation", "api characteristic"]
    miss_rows = []
    miss_totals = Counter()
    for data in collected:
        counts = Counter(
            r["Miss category"].lower() for r in data["rows"] if r["Origin"] == "Student-added"
        )
        miss_totals.update(counts)
        other = sum(v for k, v in counts.items() if k not in categories)
        miss_totals["other"] += other
        miss_rows.append([
            data["cfg"]["title"], "", "",
            counts.get("prompt quality", 0), counts.get("model limitation", 0),
            counts.get("api characteristic", 0), other,
        ])
    miss_rows.append([
        "TOTAL", "", "",
        miss_totals["prompt quality"], miss_totals["model limitation"],
        miss_totals["api characteristic"], miss_totals["other"],
    ])
    row = write_table(
        ws, row,
        ["API", "", "", "Prompt quality", "Model limitation", "API characteristic", "Other"],
        miss_rows, [30, 6, 46, 14, 16, 18, 10],
    )

    # --- bugs ---
    ws.cell(row=row, column=1, value="4. Bugs reported (GitHub Issues)").font = SECTION_FONT
    row += 1
    bug_headers = list(bugs[0].keys()) if bugs else ["Bug ID"]
    bug_rows = [[b.get(h, "") for h in bug_headers] for b in bugs]
    write_table(ws, row, bug_headers, bug_rows, [16, 34, 10, 12, 10, 14])

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 34
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument(
        "-o", "--output",
        default=str(TEST_CASES_DIR / f"{STUDENT_ID}_HW06_TestCases.xlsx"),
        help="output .xlsx path (default: artifacts/test-cases/<StudentID>_HW06_TestCases.xlsx)",
    )
    args = parser.parse_args()

    collected = []
    for cfg in FR_CONFIG:
        print(f"Reading {cfg['key']} ...")
        collected.append(collect_fr(cfg))

    bugs = parse_bugs(BUG_REPORT)
    print(f"Reading BUG_REPORT.md ... {len(bugs)} bugs")

    wb = Workbook()
    wb.remove(wb.active)
    for data in collected:
        write_case_sheet(wb, data)
    write_summary_sheet(wb, collected, bugs)
    wb.active = 0

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    total = sum(len(d["rows"]) for d in collected)
    print(f"\nWrote {out} — {len(collected) + 1} sheets, {total} test cases, {len(bugs)} bugs")
    return 0


if __name__ == "__main__":
    sys.exit(main())
